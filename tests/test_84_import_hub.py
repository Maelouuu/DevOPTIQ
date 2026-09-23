# -*- coding: utf-8 -*-
"""La fenêtre « Importer des données » (`/api/import`), pages Carte et Comptes.

Page Carte : trois natures (rôles, tâches, outils) et un import multiple.
Page Comptes : les collaborateurs, et eux seuls. Un même parcours pour tous :
lire → (organiser avec l'IA) → vérifier → importer.

Ce que ces tests tiennent, dans l'ordre où l'utilisateur le rencontre :
* un fichier au bon format est lu TEL QUEL — en-tête plus bas, anglais, csv
  au point-virgule, cellules fusionnées ;
* un fichier qui ne l'est pas n'est PAS deviné par le code : c'est l'IA qui
  désigne les colonnes, et elle ne peut rien écrire d'autre que des numéros ;
* une liste de COLLABORATEURS déposée sur la carte n'y est jamais importée —
  ni lue comme des rôles : elle est repérée et renvoyée à la page Comptes,
  où elle est au contraire lue comme ce qu'elle est ;
* le statut de chaque ligne dépend de la PORTÉE, et une carto où le compte
  n'écrit pas n'en fait jamais partie ;
* l'import revérifie tout, écrit en une transaction, et un rôle importé
  survit à l'enregistrement de la carte (`hors_carte`).
"""
import io
import json

import openpyxl
import pytest

from Code.extensions import db

API = "/api/import"


# ══════════════════════════════════════════════════════════════════════
#  Décor
# ══════════════════════════════════════════════════════════════════════
@pytest.fixture(autouse=True)
def _session_rendue(client):
    """`client` est partagé par tous les fichiers : on rend la session telle
    qu'on l'a trouvée, sinon les suivants tourneraient sous un autre compte."""
    with client.session_transaction() as s:
        avant = dict(s)
    yield
    with client.session_transaction() as s:
        s.clear()
        s.update(avant)


@pytest.fixture()
def scene(app):
    """Un administrateur et un coordinateur, deux cartos à l'administrateur
    (A et B), une carto à quelqu'un d'autre (C, hors de portée), un compte
    qui n'écrit dans aucune carto."""
    from Code.models.models import (Activities, Entity, Role, Task, Tool, User, UserRole,
                                    Competency, activity_roles, task_roles)
    from Code.security import hash_password
    with app.app_context():
        admin = User(first_name="T84", last_name="Admin", email="t84.admin@devoptiq.com",
                     password=hash_password("Motdepasse123!"), status="admin")
        coord = User(first_name="T84", last_name="Coord", email="t84.coord@devoptiq.com",
                     password=hash_password("Motdepasse123!"), status="coordinateur")
        autre = User(first_name="T84", last_name="Autre", email="t84.autre@devoptiq.com",
                     password=hash_password("Motdepasse123!"), status="user")
        lecteur = User(first_name="T84", last_name="Lecteur", email="t84.lecteur@devoptiq.com",
                       password=hash_password("Motdepasse123!"), status="user")
        db.session.add_all([admin, coord, autre, lecteur])
        db.session.flush()
        a = Entity(name="T84 Carto A", owner_id=admin.id)
        b = Entity(name="T84 Carto B", owner_id=admin.id)
        c = Entity(name="T84 Carto C", owner_id=autre.id)
        k = Entity(name="T84 Carto Coord", owner_id=coord.id)
        db.session.add_all([a, b, c, k])
        db.session.flush()
        db.session.add_all([
            Activities(name="Chiffrer l'offre", entity_id=a.id, shape_id="t84_a1"),
            Activities(name="Qualifier le besoin client", entity_id=a.id, shape_id="t84_a2"),
            Activities(name="Chiffrer l'offre", entity_id=b.id, shape_id="t84_b1"),
            Activities(name="Chiffrer l'offre", entity_id=c.id, shape_id="t84_c1"),
            Role(name="Qualité", name_fr="Qualité", name_en="Quality", entity_id=a.id),
        ])
        db.session.commit()
        ids = {"admin": admin.id, "coord": coord.id, "autre": autre.id, "lecteur": lecteur.id,
               "a": a.id, "b": b.id, "c": c.id, "k": k.id}
    yield ids
    with app.app_context():
        ents = [ids["a"], ids["b"], ids["c"], ids["k"]]
        acts = [x.id for x in Activities.query.filter(Activities.entity_id.in_(ents)).all()]
        taches = [x.id for x in Task.query.filter(Task.activity_id.in_(acts or [-1])).all()]
        roles = [x.id for x in Role.query.filter(Role.entity_id.in_(ents)).all()]
        db.session.execute(task_roles.delete().where(task_roles.c.task_id.in_(taches or [-1])))
        db.session.execute(activity_roles.delete().where(
            activity_roles.c.activity_id.in_(acts or [-1])))
        for tache in Task.query.filter(Task.id.in_(taches or [-1])).all():
            tache.tools = []
        db.session.flush()
        Task.query.filter(Task.id.in_(taches or [-1])).delete(synchronize_session=False)
        Competency.query.filter(Competency.activity_id.in_(acts or [-1])).delete(
            synchronize_session=False)
        Activities.query.filter(Activities.id.in_(acts or [-1])).delete(synchronize_session=False)
        Tool.query.filter(Tool.entity_id.in_(ents)).delete(synchronize_session=False)
        UserRole.query.filter(UserRole.role_id.in_(roles or [-1])).delete(synchronize_session=False)
        Role.query.filter(Role.id.in_(roles or [-1])).delete(synchronize_session=False)
        Entity.query.filter(Entity.id.in_(ents)).delete(synchronize_session=False)
        User.query.filter(User.email.like("t84%")).delete(synchronize_session=False)
        db.session.commit()


def _connecte(client, app, uid, entity_id=None, lang="fr"):
    from Code.models.models import User
    with app.app_context():
        mail = db.session.get(User, uid).email
    with client.session_transaction() as s:
        s.clear()
        s["user_id"] = uid
        s["user_email"] = mail
        s["lang"] = lang
        if entity_id:
            s["active_entity_id"] = entity_id


def _xlsx(*feuilles):
    """[(nom, [lignes])] → octets d'un classeur."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nom, lignes in feuilles:
        ws = wb.create_sheet(nom)
        for l in lignes:
            ws.append(l)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _lire(client, type_, fichiers, **extra):
    """fichiers : [(nom, octets)]."""
    data = {"type": type_, **extra}
    data["fichiers"] = [(io.BytesIO(o), n) for n, o in fichiers]
    return client.post(f"{API}/lire", data=data, content_type="multipart/form-data")


def _verifier(client, type_, lignes, cibles, **extra):
    return client.post(f"{API}/verifier", data=json.dumps(
        {"type": type_, "lignes": lignes, "cibles": cibles, **extra}),
        content_type="application/json")


def _importer(client, parts, cibles, **extra):
    return client.post(f"{API}/importer",
                       data=json.dumps({"parts": parts, "cibles": cibles, **extra}),
                       content_type="application/json")


# ══════════════════════════════════════════════════════════════════════
#  1 · Lire un fichier tel qu'il est
# ══════════════════════════════════════════════════════════════════════
class TestLectureTelleQuelle:

    def test_un_fichier_au_format_attendu_est_lu_tel_quel(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Postes", [["Nom du rôle", "Mission"],
                              ["Acheteur", "Négocier les achats"], ["Logisticien", ""]]))
        p = _lire(client, "roles", [("roles.xlsx", f)]).get_json()["parts"][0]
        assert p["reconnu"] and p["source"] == "FICHIER"
        assert [l["nom"] for l in p["lignes"]] == ["Acheteur", "Logisticien"]
        assert p["lignes"][0]["mission"] == "Négocier les achats"

    def test_l_en_tete_peut_etre_plus_bas_et_en_anglais(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Sheet1", [["Export RH — septembre"], [None], ["Job title", "Purpose", "Site"],
                              ["Buyer", "Buys things", "Lyon"]]))
        p = _lire(client, "roles", [("x.xlsx", f)]).get_json()["parts"][0]
        assert p["reconnu"] and p["ligne_entete"] == 2
        assert p["lignes"] == [{"nom": "Buyer", "mission": "Buys things", "_i": 0}]

    def test_un_csv_au_point_virgule_encode_a_la_windows(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        octets = "Outil;Usage\r\nSAP;Gestion des achats\r\nÉtau;Serrage\r\n".encode("cp1252")
        p = _lire(client, "outils", [("outils.csv", octets)]).get_json()["parts"][0]
        assert p["reconnu"]
        assert [l["nom"] for l in p["lignes"]] == ["SAP", "Étau"]

    def test_ce_qui_n_est_pas_reconnu_n_est_pas_devine(self, app, client, scene):
        """Pas de colonne reconnaissable : le code ne tente rien — c'est le
        rôle de l'IA, dont la lecture est montrée avant d'être acceptée."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Moyen", "Commentaire"], ["Presse 250 t", "Atelier 2"]]))
        p = _lire(client, "outils", [("o.xlsx", f)]).get_json()["parts"][0]
        assert not p["reconnu"] and p["manquants"] == ["nom"]
        assert p["lignes"] == []
        assert p["echantillon"][1] == ["Presse 250 t", "Atelier 2"]

    def test_les_cellules_fusionnees_propagent_l_activite(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Tâches", [["Activité", "Tâche", "Garant"],
                             ["Chiffrer l'offre", "Lire le cahier des charges", "Chef de projet"],
                             [None, "Estimer les heures", None],
                             ["Qualifier le besoin client", None, None],
                             [None, "Appeler le client", None]]))
        p = _lire(client, "taches", [("t.xlsx", f)]).get_json()["parts"][0]
        assert [(l["activite"], l["tache"], l["garant"]) for l in p["lignes"]] == [
            ("Chiffrer l'offre", "Lire le cahier des charges", "Chef de projet"),
            ("Chiffrer l'offre", "Estimer les heures", "Chef de projet"),
            # ⚠️ le garant du bloc précédent ne déborde pas sur l'activité suivante
            ("Qualifier le besoin client", "Appeler le client", ""),
        ]

    def test_les_mentions_d_absence_disparaissent_des_la_lecture(self, app, client, scene):
        """« No Special skills required » (27 lignes du fichier client) n'est
        pas une compétence : l'aperçu montre ce qui sera vraiment importé."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("T", [["Activity", "Task", "Tools", "Skills"],
                         ["Chiffrer l'offre", "Estimer", "SAP; - ; Excel",
                          "No Special skills required"]]))
        l = _lire(client, "taches", [("t.xlsx", f)]).get_json()["parts"][0]["lignes"][0]
        assert l["outils"] == "SAP, Excel"
        assert l["competences"] == ""

    def test_un_format_non_pris_en_charge_est_refuse_proprement(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        d = _lire(client, "roles", [("notes.txt", b"bonjour")]).get_json()
        assert d["parts"] == [] and d["erreurs"][0]["fichier"] == "notes.txt"


# ══════════════════════════════════════════════════════════════════════
#  2 · Les collaborateurs ne s'importent pas ici
# ══════════════════════════════════════════════════════════════════════
class TestLesCollaborateursRestentALaPageComptes:

    @pytest.mark.parametrize("entetes, valeurs", [
        (["Prénom", "Nom", "Poste"], ["Jean", "Dupont", "Acheteur"]),
        (["Nom", "Contact"], ["Jean Dupont", "jean.dupont@exemple.fr"]),
        (["Collaborateur", "E-mail"], ["DUPONT Jean", "jean.dupont@exemple.fr"]),
    ])
    def test_une_liste_de_personnes_n_est_pas_lue_comme_des_roles(
            self, app, client, scene, entetes, valeurs):
        """⚠️ Sans ce repérage, « Nom » passait pour le nom d'un rôle : chaque
        nom de famille devenait un rôle de la carto."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [entetes, valeurs, valeurs]))
        p = _lire(client, "roles", [("personnes.xlsx", f)]).get_json()["parts"][0]
        assert p["type"] == "comptes" and not p["reconnu"] and p["lignes"] == []

    def test_un_role_au_nom_compose_n_est_pas_une_personne(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Nom", "Mission"], ["Chef de projet", "Piloter"],
                              ["Responsable qualité", "Garantir"]]))
        p = _lire(client, "roles", [("r.xlsx", f)]).get_json()["parts"][0]
        assert p["type"] == "roles" and p["reconnu"]

    def test_un_tableau_de_taches_complet_l_emporte(self, app, client, scene):
        """Une colonne d'e-mails dans un tableau de tâches ne fait pas de lui
        une liste de personnes : l'activité et la tâche sont là."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Tâches", [["Activité", "Tâche", "Nom", "Mail"],
                              ["Chiffrer l'offre", "Estimer", "Dupont", "j.dupont@x.fr"],
                              ["Chiffrer l'offre", "Relire", "Martin", "l.martin@x.fr"]]))
        p = _lire(client, "taches", [("t.xlsx", f)]).get_json()["parts"][0]
        assert p["type"] == "taches" and p["reconnu"]

    def test_l_utilisateur_peut_passer_outre(self, app, client, scene):
        """Le repérage peut se tromper : dire soi-même ce que contient la
        feuille suffit à la faire lire comme telle."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Nom", "Contact"], ["Chef de projet", "pmo@x.fr"],
                              ["Acheteur", "achats@x.fr"]]))
        assert _lire(client, "roles", [("r.xlsx", f)]).get_json()["parts"][0]["type"] == "comptes"
        p = _lire(client, "roles", [("r.xlsx", f)], comme="roles").get_json()["parts"][0]
        assert p["type"] == "roles" and [l["nom"] for l in p["lignes"]] == ["Chef de projet", "Acheteur"]

    def test_la_carte_ne_propose_pas_les_comptes(self, app, client, scene):
        """Sur la carte, la question ne se pose pas : pas de carte
        « Collaborateurs », pas de feuille de comptes dans un import multiple —
        mais un renvoi vers la page qui, elle, sait les créer."""
        _connecte(client, app, scene["admin"], scene["a"])
        assert _importer(client, [{"type": "comptes", "lignes": []}], [scene["a"]]).status_code == 400
        ctx = client.get(f"{API}/contexte").get_json()
        assert set(ctx["natures"]) == {"roles", "taches", "outils", "multiple"}
        assert ctx["pour"] == "carto"
        # la fenêtre sait où renvoyer — et si le lien mène quelque part
        assert ctx["comptes"]["url"].startswith("/comptes") and ctx["comptes"]["peut"] is True


# ══════════════════════════════════════════════════════════════════════
#  2 bis · Les comptes s'importent, mais depuis la page Comptes
# ══════════════════════════════════════════════════════════════════════
class TestLesComptesSImportentDepuisLaPageComptes:

    def test_le_contexte_de_la_page_comptes_n_offre_que_les_comptes(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        ctx = client.get(f"{API}/contexte?pour=comptes").get_json()
        assert set(ctx["natures"]) == {"users"} and ctx["pour"] == "comptes"
        assert ctx["peut"] is True
        # rien à renvoyer ailleurs : on y est
        assert "comptes" not in ctx

    def test_une_liste_de_personnes_y_est_lue_comme_telle(self, app, client, scene):
        """Le repérage qui l'écarte de la carte ne doit pas l'écarter ICI."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Prénom", "Nom", "E-mail"],
                              ["Jean", "Dupont", "t84.jean@x.fr"]]))
        p = _lire(client, "users", [("u.xlsx", f)]).get_json()["parts"][0]
        assert p["type"] == "users" and p["reconnu"]
        assert [l["email"] for l in p["lignes"]] == ["t84.jean@x.fr"]

    def test_un_nom_complet_se_scinde_selon_les_capitales(self, app, client, scene):
        """Cas courant d'un export RH : « DUPONT Jean » dans une seule colonne."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Effectif", [["Nom complet", "Mail"], ["DUPONT Jean", "jean@x.fr"],
                                ["Marie CURIE", "marie@x.fr"]]))
        p = _lire(client, "users", [("u.xlsx", f)]).get_json()["parts"][0]
        assert p["reconnu"], p["manquants"]
        assert [(l["prenom"], l["nom"]) for l in p["lignes"]] == [("Jean", "DUPONT"),
                                                                  ("Marie", "CURIE")]

    def test_les_comptes_se_verifient_sans_aucune_carto(self, app, client, scene):
        """Un compte n'appartient à aucune carto : il se vérifie et s'importe
        sans en choisir une. ⚠️ Et aucun mot de passe ne revient vers l'écran."""
        from Code.translations import t
        _connecte(client, app, scene["coord"], scene["k"])
        lignes = [
            {"prenom": "Jean", "nom": "Neuf", "email": "t84.jean@x.fr", "_i": 0},
            {"prenom": "T84", "nom": "Admin", "email": "T84.ADMIN@devoptiq.com", "_i": 1},
            {"prenom": "Sans", "nom": "Mail", "email": "pas-un-mail", "_i": 2},
            {"prenom": "Chef", "nom": "Suprême", "email": "t84.chef@x.fr",
             "statut": "Administrateur", "_i": 3},
            {"prenom": "Paul", "nom": "Flou", "email": "t84.paul@x.fr", "statut": "Stagiaire",
             "mot_de_passe": "zq9", "_i": 4},
        ]
        d = _verifier(client, "users", lignes, []).get_json()
        assert [l["statut"] for l in d["lignes"]] == [
            "nouveau", "present", "invalide", "invalide", "nouveau"]
        # un coordinateur ne crée pas d'administrateur
        assert d["lignes"][3]["raison"] == t("imph.st_statut_superieur", "fr")
        assert "Stagiaire" in d["lignes"][4]["avertissement"]
        assert all("mot_de_passe" not in l for l in d["lignes"])
        assert "zq9" not in json.dumps(d)

    def test_un_role_inconnu_est_signale_ou_annonce(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [{"prenom": "Jean", "nom": "Neuf", "email": "t84.jean@x.fr",
                   "role": "Magasinier"}]
        sans = _verifier(client, "users", lignes, [scene["a"]]).get_json()["lignes"][0]
        assert "Magasinier" in sans["avertissement"]
        avec = _verifier(client, "users", lignes, [scene["a"]],
                         options={"creer_roles": True}).get_json()["lignes"][0]
        assert "avertissement" not in avec and "Magasinier" in avec["info"]

    def test_des_comptes_avec_leur_role_et_leurs_identifiants(self, app, client, scene):
        from Code.models.models import Role, User, UserRole
        from Code.security import verify_password
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [
            {"prenom": "Jean", "nom": "Neuf", "email": "T84.Jean@X.fr", "statut": "Coordinator",
             "role": "Quality", "mot_de_passe": "Secret123", "_i": 0},
            {"prenom": "Lou", "nom": "Sans", "email": "t84.lou@x.fr", "_i": 1},
        ]
        res = _importer(client, [{"type": "users", "lignes": lignes}],
                        [scene["a"]]).get_json()["resultats"][0]
        assert res["crees"] == 2 and res["roles_attribues"] == 1
        # seul le mot de passe GÉNÉRÉ est rendu, une fois
        assert [x["email"] for x in res["identifiants"]] == ["t84.lou@x.fr"]
        with app.app_context():
            jean = User.query.filter_by(email="t84.jean@x.fr").one()
            assert jean.status == "coordinateur"
            assert verify_password(jean.password, "Secret123")
            qualite = Role.query.filter_by(entity_id=scene["a"], name="Qualité").one()
            assert UserRole.query.filter_by(user_id=jean.id, role_id=qualite.id).count() == 1
            lou = User.query.filter_by(email="t84.lou@x.fr").one()
            assert verify_password(lou.password, res["identifiants"][0]["mot_de_passe"])

    def test_le_statut_se_reverifie_a_l_import(self, app, client, scene):
        """Le navigateur ne décide de rien : un administrateur envoyé tel quel
        par un coordinateur n'est pas créé."""
        from Code.models.models import User
        _connecte(client, app, scene["coord"], scene["k"])
        res = _importer(client, [{"type": "users", "lignes": [
            {"prenom": "Chef", "nom": "Suprême", "email": "t84.chef@x.fr", "statut": "admin"}]}],
            []).get_json()["resultats"][0]
        assert res["crees"] == 0
        with app.app_context():
            assert User.query.filter_by(email="t84.chef@x.fr").count() == 0

    def test_qui_ne_cree_pas_de_comptes_n_en_importe_pas(self, app, client, scene):
        """⚠️ Écrire dans une carto ne donne PAS le droit de créer des comptes :
        c'est le droit de la page Comptes qui décide, et lui seul."""
        _connecte(client, app, scene["autre"], scene["c"])
        ctx = client.get(f"{API}/contexte?pour=comptes").get_json()
        assert ctx["peut"] is False
        # il est propriétaire de sa carto : il y importe bien des rôles
        assert client.get(f"{API}/contexte").get_json()["peut"] is True
        assert _verifier(client, "users", [{"prenom": "A", "nom": "B", "email": "t84.x@x.fr"}],
                         []).status_code == 403
        r = _importer(client, [{"type": "users", "lignes": [
            {"prenom": "A", "nom": "B", "email": "t84.x@x.fr"}]}], [])
        assert r.status_code == 403


class TestImportMultiple:

    def test_chaque_feuille_est_reconnue(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Rôles", [["Nom", "Mission"], ["Acheteur", ""]]),
                  ("Collaborateurs", [["Prénom", "Nom", "E-mail"], ["Jean", "Dupont", "j@x.fr"]]),
                  ("Outils", [["Nom", "Description"], ["SAP", "ERP"]]),
                  ("Tâches", [["Activité", "Tâche"], ["Chiffrer l'offre", "Estimer"]]),
                  ("Vide", []))
        parts = _lire(client, "multiple", [("tout.xlsx", f)]).get_json()["parts"]
        assert {p["feuille"]: (p["type"], p["reconnu"]) for p in parts} == {
            "Rôles": ("roles", True), "Collaborateurs": ("comptes", False),
            "Outils": ("outils", True), "Tâches": ("taches", True)}

    def test_une_feuille_ambigue_le_dit(self, app, client, scene):
        """« Nom | Description » : un rôle ou un outil ? Le nom de la feuille ne
        tranche pas — l'écran demande de choisir au lieu de décider en silence."""
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Nom", "Description"], ["SAP", "ERP"]]))
        p = _lire(client, "multiple", [("x.xlsx", f)]).get_json()["parts"][0]
        assert p["ambigu"] == ["roles", "outils"]

    def test_la_nature_se_corrige_et_la_feuille_est_relue(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        f = _xlsx(("Feuil1", [["Nom", "Description"], ["SAP", "ERP"]]))
        p = _lire(client, "multiple", [("x.xlsx", f)], comme="outils",
                  feuille="Feuil1").get_json()["parts"][0]
        assert p["type"] == "outils" and p["lignes"][0]["description"] == "ERP"


# ══════════════════════════════════════════════════════════════════════
#  3 · L'IA désigne les colonnes, elle n'écrit rien
# ══════════════════════════════════════════════════════════════════════
class TestOrganiserAvecLIA:

    FICHIER = _xlsx(("Feuil1", [["Bloc", "Ce qu'on fait", "Avec quoi"],
                               ["Chiffrer l'offre", "Estimer les heures", "SAP"],
                               ["Chiffrer l'offre", "Relire le devis", "Excel"]]))

    def _organiser(self, client, type_="taches"):
        return client.post(f"{API}/organiser", data={
            "type": type_, "fichier": (io.BytesIO(self.FICHIER), "t.xlsx")},
            content_type="multipart/form-data")

    def test_les_valeurs_viennent_du_fichier(self, app, client, scene, monkeypatch):
        from Code.routes import import_hub
        vu = {}

        def ia(systeme, contenu):
            vu.update(contenu)
            return {"type": "taches", "ligne_entete": 0,
                    "colonnes": {"activite": 0, "tache": 1, "outils": 2},
                    "confiance": "high", "remarque": "« Bloc » porte l'activité."}
        monkeypatch.setattr(import_hub, "_appeler_ia", ia)
        _connecte(client, app, scene["admin"], scene["a"])
        p = self._organiser(client).get_json()["part"]
        assert p["reconnu"] and p["source"] == "IA" and p["confiance"] == "high"
        assert [(l["activite"], l["tache"], l["outils"]) for l in p["lignes"]] == [
            ("Chiffrer l'offre", "Estimer les heures", "SAP"),
            ("Chiffrer l'offre", "Relire le devis", "Excel")]
        # ce que l'IA a vu : les premières lignes, cellules numérotées, et le schéma
        assert vu["lignes"][1]["cellules"] == {"0": "Chiffrer l'offre", "1": "Estimer les heures",
                                               "2": "SAP"}
        assert {c["champ"] for c in vu["schemas"]["taches"]} >= {"activite", "tache"}

    def test_l_ia_ne_peut_pas_glisser_de_lignes(self, app, client, scene, monkeypatch):
        """Une réponse qui porterait des DONNÉES est ignorée : seuls les numéros
        de colonnes comptent, les valeurs sont relues dans le fichier."""
        from Code.routes import import_hub
        monkeypatch.setattr(import_hub, "_appeler_ia", lambda s, c: {
            "type": "taches", "ligne_entete": 0, "colonnes": {"activite": 0, "tache": 1},
            "lignes": [{"activite": "Inventée", "tache": "Tâche fantôme"}]})
        _connecte(client, app, scene["admin"], scene["a"])
        p = self._organiser(client).get_json()["part"]
        assert "fantôme" not in json.dumps(p, ensure_ascii=False)
        assert [l["tache"] for l in p["lignes"]] == ["Estimer les heures", "Relire le devis"]

    def test_une_colonne_hors_du_fichier_est_ecartee(self, app, client, scene, monkeypatch):
        from Code.routes import import_hub
        monkeypatch.setattr(import_hub, "_appeler_ia", lambda s, c: {
            "type": "taches", "ligne_entete": 0, "colonnes": {"activite": 42, "tache": -1}})
        _connecte(client, app, scene["admin"], scene["a"])
        p = self._organiser(client).get_json()["part"]
        assert not p["reconnu"]
        assert set(p["manquants"]) == {"activite", "tache"}

    def test_sans_ia_un_refus_explicite(self, app, client, scene, monkeypatch):
        from Code.routes import import_hub
        from Code.translations import t
        monkeypatch.setattr(import_hub, "_appeler_ia", lambda s, c: None)
        _connecte(client, app, scene["admin"], scene["a"])
        r = self._organiser(client)
        assert r.status_code == 503
        assert r.get_json()["error"] == t("imph.err_ia", "fr")

    def test_rapprocher_propose_sans_appliquer(self, app, client, scene, monkeypatch):
        """L'IA rapproche une activité du fichier d'une activité de la carto par
        le SENS ; elle rend une proposition, jamais un rattachement."""
        from Code.routes import import_hub
        monkeypatch.setattr(import_hub, "_appeler_ia", lambda s, c: {"resolved": [
            {"activity_name_excel": "Price the bid", "activity_id": [
                a["name"] for a in c["db_activities"]].index("Chiffrer l'offre"),
             "confidence": "high", "match_reason": "même sens"},
            {"activity_name_excel": "Price the bid", "activity_id": 999},
            {"activity_name_excel": "Pas demandé", "activity_id": 0}]})
        _connecte(client, app, scene["admin"], scene["a"])
        r = client.post(f"{API}/rapprocher", data=json.dumps(
            {"noms": ["Price the bid"], "cibles": [scene["a"]]}), content_type="application/json")
        assert r.get_json()["propositions"] == {"Price the bid": {
            "activite": "Chiffrer l'offre", "confiance": "high", "raison": "même sens"}}

    @pytest.mark.parametrize("lang, langue", [("fr", "français"), ("en", "English")])
    def test_la_justification_est_ecrite_dans_la_langue_de_l_ecran(
            self, app, client, scene, monkeypatch, lang, langue):
        """La raison de chaque rapprochement s'affiche telle quelle dans le
        compte rendu : l'IA doit l'écrire dans la langue de celui qui le lit."""
        from Code.routes import import_hub
        vu = {}
        monkeypatch.setattr(import_hub, "_appeler_ia",
                            lambda s, c: vu.update(c) or {"resolved": []})
        _connecte(client, app, scene["admin"], scene["a"], lang=lang)
        client.post(f"{API}/rapprocher", data=json.dumps(
            {"noms": ["Price the bid"], "cibles": [scene["a"]]}), content_type="application/json")
        assert vu["langue_des_remarques"] == langue
        _connecte(client, app, scene["admin"], scene["a"])


# ══════════════════════════════════════════════════════════════════════
#  4 · Vérifier : le statut dépend de la portée
# ══════════════════════════════════════════════════════════════════════
class TestVerifier:

    def test_le_statut_d_un_role_ne_depend_pas_des_cartos_visees(self, app, client, scene):
        """Un rôle appartient à l'entreprise : il existe ou il n'existe pas."""
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [{"nom": "Qualité", "_i": 0}, {"nom": "Logistique", "_i": 1},
                  {"nom": "", "_i": 2}, {"nom": "QUALITE", "_i": 3}]
        une = _verifier(client, "roles", lignes, [scene["a"]]).get_json()["lignes"]
        assert [l["statut"] for l in une] == ["present", "nouveau", "invalide", "invalide"]
        deux = _verifier(client, "roles", lignes, [scene["a"], scene["b"]]).get_json()["lignes"]
        assert [l["statut"] for l in deux] == [l["statut"] for l in une]

    def test_la_traduction_d_un_role_compte_comme_le_role(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        l = _verifier(client, "roles", [{"nom": "Quality"}], [scene["a"]]).get_json()["lignes"][0]
        assert l["statut"] == "present"

    def test_une_carto_ou_l_on_n_ecrit_pas_n_est_jamais_visee(self, app, client, scene):
        """La carto C appartient à un autre compte : son id envoyé par le
        navigateur est simplement ignoré."""
        _connecte(client, app, scene["admin"], scene["a"])
        r = _verifier(client, "roles", [{"nom": "X"}], [scene["c"]])
        assert r.status_code == 400
        r = _verifier(client, "roles", [{"nom": "X"}], [scene["a"], scene["c"]])
        assert r.get_json()["cibles"] == [scene["a"]]

    def test_les_taches_se_rattachent_a_leur_activite(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [
            {"activite": "Chiffrer l'offre", "tache": "Estimer", "_i": 0},
            {"activite": "CHIFFRER L'OFFRE !", "tache": "Relire", "_i": 1},
            {"activite": "Qualifier le besoin du client", "tache": "Appeler", "_i": 2},
            {"activite": "Totalement autre chose", "tache": "Rien", "_i": 3},
        ]
        d = _verifier(client, "taches", lignes, [scene["a"], scene["b"]]).get_json()
        modes = {g["activite_fichier"]: (g["mode"], g["choix"]) for g in d["groupes"]}
        # la casse et la ponctuation ne séparent pas deux lignes de la même activité
        assert len(d["groupes"]) == 3
        assert [l["tache"] for l in d["groupes"][0]["lignes"]] == ["Estimer", "Relire"]
        assert modes["Chiffrer l'offre"] == ("exact", "Chiffrer l'offre")
        assert modes["Qualifier le besoin du client"] == ("proche", "Qualifier le besoin client")
        assert modes["Totalement autre chose"] == ("a_rattacher", None)
        # « Chiffrer l'offre » existe dans les DEUX cartos
        g = next(g for g in d["groupes"] if g["activite_fichier"] == "Chiffrer l'offre")
        assert g["n_cartos"] == 2 and g["n_cibles"] == 2
        assert d["totaux"]["a_rattacher"] == 1

    def test_le_choix_de_l_utilisateur_l_emporte(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [{"activite": "Price the bid", "tache": "Estimer", "_i": 0},
                  {"activite": "Chiffrer l'offre", "tache": "Relire", "_i": 1}]
        d = _verifier(client, "taches", lignes, [scene["a"]],
                      choix={"Price the bid": "Chiffrer l'offre", "Chiffrer l'offre": ""}).get_json()
        modes = {g["activite_fichier"]: g["mode"] for g in d["groupes"]}
        assert modes == {"Price the bid": "manuel", "Chiffrer l'offre": "ignore"}
        assert [l["statut"] for l in d["lignes"]] == ["nouveau", "ignore"]


# ══════════════════════════════════════════════════════════════════════
#  5 · Importer
# ══════════════════════════════════════════════════════════════════════
class TestImporter:

    def test_des_roles_dans_deux_cartos_hors_de_la_carte(self, app, client, scene):
        from Code.models.models import Role
        _connecte(client, app, scene["admin"], scene["a"])
        r = _importer(client, [{"type": "roles", "lignes": [
            {"nom": "Logistique", "mission": "Livrer à l'heure", "_i": 0},
            {"nom": "Qualité", "_i": 1}]}], [scene["a"], scene["b"]])
        res = r.get_json()["resultats"][0]
        # Un seul « Logistique » pour l'entreprise ; « Qualité » existait déjà.
        assert res["crees"] == 1
        with app.app_context():
            logs = Role.query.filter_by(name="Logistique").all()
            assert len(logs) == 1 and logs[0].hors_carte
            assert logs[0].mission_generale == "Livrer à l'heure"

    def test_un_role_importe_survit_a_l_enregistrement_de_la_carte(self, app, client, scene):
        """Vérifié ROUGE sans `hors_carte` : `_sync_carto_to_db` effaçait tout
        rôle absent des bandes — importer des rôles ne servait à rien."""
        from Code.models.models import Entity, Role
        from Code.routes.cartography_editor import _sync_carto_to_db
        _connecte(client, app, scene["coord"], scene["k"])
        _importer(client, [{"type": "roles", "lignes": [{"nom": "Importé T84"}]}], [scene["k"]])
        with app.app_context():
            ent = db.session.get(Entity, scene["k"])
            _sync_carto_to_db(ent, {"shapes": [], "connections": [],
                                    "bands": [{"id": "b1", "label": "Bande T84", "y": 0, "h": 200}]})
            db.session.commit()
            noms = {r.name for r in Role.query.filter_by(entity_id=scene["k"]).all()}
        assert {"Importé T84", "Bande T84"} <= noms

    def test_des_taches_avec_leurs_outils_roles_et_competences(self, app, client, scene):
        from Code.models.models import Activities, Competency, Task, Tool
        _connecte(client, app, scene["admin"], scene["a"])
        lignes = [{"activite": "Chiffrer l'offre", "tache": "Estimer les heures",
                   "outils": "SAP, Excel", "realisateur": "Commercial",
                   "competences": "Négocier", "garant": "Chef de projet", "_i": 0}]
        res = _importer(client, [{"type": "taches", "lignes": lignes}],
                        [scene["a"], scene["b"]]).get_json()["resultats"][0]
        assert res["tasks_created"] == 2 and res["activities_updated"] == 2
        with app.app_context():
            for e in (scene["a"], scene["b"]):
                act = Activities.query.filter_by(entity_id=e, name="Chiffrer l'offre").first()
                tache = Task.query.filter_by(activity_id=act.id, name="Estimer les heures").one()
                assert {x.name for x in tache.tools} == {"SAP", "Excel"}
                assert Tool.query.filter_by(entity_id=e, name="SAP").count() == 1
                assert Competency.query.filter_by(activity_id=act.id, description="Négocier").count() == 1
        # rejouer ne double rien
        res = _importer(client, [{"type": "taches", "lignes": lignes}],
                        [scene["a"], scene["b"]]).get_json()["resultats"][0]
        assert res["tasks_created"] == 0

    def test_une_tache_ne_va_que_la_ou_son_activite_existe(self, app, client, scene):
        from Code.models.models import Activities, Task
        _connecte(client, app, scene["admin"], scene["a"])
        _importer(client, [{"type": "taches", "lignes": [
            {"activite": "Qualifier le besoin client", "tache": "Appeler T84", "_i": 0}]}],
            [scene["a"], scene["b"]])
        with app.app_context():
            n = Task.query.join(Activities).filter(Task.name == "Appeler T84").count()
        assert n == 1

    def test_un_import_multiple_ordonne_ses_ecritures(self, app, client, scene):
        """L'outil décrit par la feuille « Outils » existe AVANT que la feuille
        « Tâches » ne s'en serve — même envoyées dans l'autre sens : la tâche
        retrouve l'outil avec sa description au lieu d'en créer un vide."""
        from Code.models.models import Tool
        _connecte(client, app, scene["admin"], scene["a"])
        d = _importer(client, [
            {"type": "taches", "lignes": [{"activite": "Chiffrer l'offre", "tache": "Chiffrer T84",
                                           "outils": "Chiffreur T84"}]},
            {"type": "outils", "lignes": [{"nom": "Chiffreur T84", "description": "Logiciel de devis"}]},
        ], [scene["a"]]).get_json()
        assert [x["type"] for x in d["resultats"]] == ["outils", "taches"]
        with app.app_context():
            outils = Tool.query.filter_by(entity_id=scene["a"], name="Chiffreur T84").all()
            assert len(outils) == 1 and outils[0].description == "Logiciel de devis"

    def test_une_part_qui_echoue_n_en_laisse_aucune(self, app, client, scene, monkeypatch):
        from Code.models.models import Role
        from Code.routes import import_hub

        def casse(*a, **k):
            raise RuntimeError("panne")
        monkeypatch.setattr(import_hub, "_importer_taches", casse)
        _connecte(client, app, scene["admin"], scene["a"])
        r = _importer(client, [
            {"type": "roles", "lignes": [{"nom": "Fantôme T84"}]},
            {"type": "taches", "lignes": [{"activite": "Chiffrer l'offre", "tache": "X"}]},
        ], [scene["a"]])
        assert r.status_code == 500
        with app.app_context():
            assert Role.query.filter_by(name="Fantôme T84").count() == 0


class TestLesRolesDejaEnBase:
    """Les rôles créés hors de la carte AVANT ce correctif ne portaient pas la
    marque : la reprise les protège, une fois, sans toucher aux bandes."""

    def test_la_reprise_ne_marque_que_ce_qui_n_est_pas_une_bande(self, app, scene):
        from Code.models.models import Entity, Role
        from Code.roles_permanents import reprendre_roles_hors_carte
        from Code.routes.cartography_editor import _sync_carto_to_db
        with app.app_context():
            ent = db.session.get(Entity, scene["k"])
            ent.optiqcarto_data = json.dumps({"shapes": [], "connections": [], "bands": [
                {"id": "b1", "label": "Bande T84", "y": 0, "h": 200}]})
            db.session.add_all([Role(name="Bande T84", entity_id=ent.id),
                                Role(name="Créé depuis la page RH T84", entity_id=ent.id),
                                Role(name="Développeur de compétences", entity_id=ent.id)])
            sans_carte = Entity(name="T84 sans diagramme", owner_id=scene["coord"])
            db.session.add(sans_carte)
            db.session.flush()
            db.session.add(Role(name="Rôle d'une carto illisible T84", entity_id=sans_carte.id))
            db.session.commit()
            try:
                assert reprendre_roles_hors_carte(force=True,
                                                  entity_ids=[ent.id, sans_carte.id]) == 1
                marque = {r.name: r.hors_carte for r in Role.query.filter(
                    Role.entity_id.in_([ent.id, sans_carte.id])).all()}
                assert marque["Créé depuis la page RH T84"] is True
                assert marque["Bande T84"] is False
                assert marque["Rôle d'une carto illisible T84"] is False
                # ⚠️ une seule fois : le marqueur est posé, la reprise ne se rejoue pas
                assert reprendre_roles_hors_carte() == 0
                # et le rôle protégé survit bien à l'enregistrement de la carte
                _sync_carto_to_db(ent, json.loads(ent.optiqcarto_data))
                db.session.commit()
                noms = {r.name for r in Role.query.filter_by(entity_id=ent.id).all()}
                assert "Créé depuis la page RH T84" in noms
            finally:
                Role.query.filter_by(entity_id=sans_carte.id).delete()
                db.session.delete(sans_carte)
                db.session.commit()


class TestLAncienneRouteDImport:

    def test_ecrire_demande_le_droit_d_ecrire(self, app, client, scene):
        """`/api/import-full/inject` écrivait dans l'entité active sans rien
        demander d'autre. Vérifié ROUGE avant ce correctif (201)."""
        from Code.models.models import Activities
        with app.app_context():
            act = Activities.query.filter_by(entity_id=scene["a"]).first().id
        _connecte(client, app, scene["autre"], scene["a"])
        r = client.post("/api/import-full/inject", data=json.dumps({"groups": [
            {"activity_id": act, "tasks": [{"name": "Intrusion T84"}]}]}),
            content_type="application/json")
        assert r.status_code == 403


# ══════════════════════════════════════════════════════════════════════
#  6 · Droits
# ══════════════════════════════════════════════════════════════════════
class TestDroits:

    def test_anonyme(self, client):
        with client.session_transaction() as s:
            s.clear()
        assert client.get(f"{API}/contexte").status_code == 401
        assert _verifier(client, "roles", [{"nom": "X"}], [1]).status_code == 403

    def test_sans_carto_ou_ecrire_rien_ne_s_importe(self, app, client, scene, monkeypatch):
        """⚠️ Les cartos sans propriétaire laissées par d'autres fichiers de
        tests sont modifiables par tous : on fixe donc le périmètre ici."""
        from Code.routes import import_hub
        monkeypatch.setattr(import_hub, "_cibles_possibles", lambda moi: [])
        _connecte(client, app, scene["lecteur"])
        ctx = client.get(f"{API}/contexte").get_json()
        assert ctx["peut"] is False and ctx["comptes"]["peut"] is False
        f = _xlsx(("R", [["Nom du rôle"], ["Acheteur"]]))
        assert _lire(client, "roles", [("r.xlsx", f)]).status_code == 403
        assert _verifier(client, "roles", [{"nom": "X"}], [scene["a"]]).status_code == 403

    def test_le_contexte_parle_la_langue_de_l_ecran(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"], lang="en")
        ctx = client.get(f"{API}/contexte").get_json()
        assert ctx["natures"]["taches"]["nom"] == "Tasks"
        assert ctx["natures"]["roles"]["champs"][0]["label"] == "Role name"
        assert ctx["active"]["id"] == scene["a"]
        assert {c["id"] for c in ctx["cibles"]} >= {scene["a"], scene["b"]}
        assert scene["c"] not in {c["id"] for c in ctx["cibles"]}


# ══════════════════════════════════════════════════════════════════════
#  7 · Le modèle téléchargé se relit tel quel, dans les deux langues
# ══════════════════════════════════════════════════════════════════════
class TestLeModeleSeRelit:

    EXEMPLES = {"roles": ["Acheteur", "Négocier"],
                "taches": ["Chiffrer l'offre", "Estimer", "", "", "", "", "", ""],
                "outils": ["SAP", "ERP"],
                "users": ["Jean", "Dupont", "t84.modele@x.fr", "", "", ""]}

    @pytest.mark.parametrize("lang", ["fr", "en"])
    def test_chaque_modele_est_reconnu(self, app, client, scene, lang):
        _connecte(client, app, scene["admin"], scene["a"], lang=lang)
        wb = openpyxl.load_workbook(io.BytesIO(client.get(f"{API}/modele/multiple").data))
        assert len(wb.worksheets) == 3
        for ws, ty in zip(wb.worksheets, ("roles", "outils", "taches")):
            ws.append(self.EXEMPLES[ty])
        out = io.BytesIO()
        wb.save(out)
        parts = _lire(client, "multiple", [("modele.xlsx", out.getvalue())]).get_json()["parts"]
        assert [(p["type"], p["reconnu"]) for p in parts] == [
            ("roles", True), ("outils", True), ("taches", True)]
        # et un modèle SEUL se relit dans son propre import — comptes compris,
        # même si l'import multiple, lui, ne les prend jamais
        for ty in ("roles", "taches", "outils", "users"):
            wb = openpyxl.load_workbook(io.BytesIO(client.get(f"{API}/modele/{ty}").data))
            wb.active.append(self.EXEMPLES[ty])
            out = io.BytesIO()
            wb.save(out)
            p = _lire(client, ty, [("m.xlsx", out.getvalue())]).get_json()["parts"][0]
            assert p["reconnu"] and p["source"] == "FICHIER", (ty, p["manquants"])


# ══════════════════════════════════════════════════════════════════════
#  8 · La fenêtre elle-même
# ══════════════════════════════════════════════════════════════════════
class TestLaFenetre:

    def test_la_page_carte_ouvre_la_nouvelle_fenetre(self, app, client, scene):
        _connecte(client, app, scene["admin"], scene["a"])
        html = client.get("/activities/map").get_data(as_text=True)
        assert 'id="imh"' in html and "js/import_hub.js" in html
        assert "import-full-overlay" not in html

    def test_les_libelles_construits_existent_dans_les_deux_langues(self):
        """Nature, description, portée et champs sont appelés par une clé
        CONSTRUITE (`imph.t_` + nature) : aucun contrôle de catalogue ne les
        voit passer."""
        from Code.routes.import_hub import TYPES
        from Code.translations import TRANSLATIONS
        cles = [f"imph.{p}_{ty}" for ty in list(TYPES) + ["multiple"] for p in ("t", "d", "p")]
        cles += [f"imph.f_{ty}_{c['cle']}" for ty, v in TYPES.items() for c in v["champs"]]
        for lang in ("fr", "en"):
            manque = [c for c in cles if c not in TRANSLATIONS[lang]]
            assert not manque, (lang, manque)

    def test_le_prompt_existe(self):
        from Code.prompts.catalog import PROMPTS
        assert "import.correspondance" in PROMPTS and "import.enrich" in PROMPTS
