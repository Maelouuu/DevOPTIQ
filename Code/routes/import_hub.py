# -*- coding: utf-8 -*-
"""Importer des données depuis un fichier — la fenêtre d'import de la page Carte.

Une carte par nature de donnée (rôles, tâches, outils) et un import MULTIPLE :
plusieurs fichiers, ou un classeur dont chaque feuille porte une nature
différente. Pour chacune, le même parcours :

  1. LIRE — le fichier est lu tel quel quand ses colonnes sont reconnues :
     synonymes français et anglais, casse et accents ignorés, ligne d'en-tête
     cherchée dans les dix premières lignes de chaque feuille.
  2. ORGANISER — sinon, l'IA désigne les colonnes : la ligne d'en-tête et, pour
     chaque champ, la colonne qui le porte. ⚠️ Elle n'écrit AUCUNE donnée : les
     lignes restent celles du fichier, relues par le code. Une IA qui
     réécrirait les lignes pourrait en inventer ; une IA qui désigne des
     colonnes se vérifie d'un coup d'œil — et l'écran montre sa lecture AVANT
     qu'on l'accepte.
  3. VÉRIFIER — chaque ligne reçoit son statut contre la base, selon la PORTÉE
     choisie. Rien n'est écrit.
  4. IMPORTER — seulement ce que l'utilisateur a gardé, en UNE transaction.

⚠️ Les COMPTES n'entrent pas ici : créer des collaborateurs reste l'affaire de
la page Comptes, réservée à ceux qui en ont le droit. Une feuille de
collaborateurs est pourtant REPÉRÉE (`COMPTES`) — pour être écartée avec un
renvoi vers la bonne page : sans ce repérage, « Prénom | Nom | E-mail »
passerait pour une liste de rôles, et chaque nom de famille deviendrait un rôle.

⚠️ La portée dépend de la donnée : un rôle, un outil se posent dans une ou
plusieurs cartos ; une tâche appartient à une activité, donc aux cartos où
cette activité existe. Et on n'écrit jamais que dans une carto où le compte
écrit déjà (`can_edit`) : un identifiant venu du navigateur ne suffit pas.
"""
import csv
import io
import json
import os
import re
import unicodedata
from difflib import SequenceMatcher

import openpyxl
from flask import Blueprint, current_app, jsonify, request, send_file, session, url_for

from Code.extensions import db
from Code.models.models import Activities, Entity, Role, Task, Tool
from Code.translations import t

import_hub_bp = Blueprint("import_hub", __name__, url_prefix="/api/import")

MAX_OCTETS = 5 * 1024 * 1024
MAX_LIGNES = 2000
MAX_FICHIERS = 8
EXTENSIONS = (".xlsx", ".xlsm", ".csv")
# L'ordre d'écriture d'un import multiple : les outils avant les tâches qui
# s'en servent — une tâche retrouve alors l'outil AVEC sa description.
ORDRE = ("roles", "outils", "taches")

# ── Le catalogue des natures ────────────────────────────────────────────────
# `aide` décrit le champ à l'IA (jamais affiché : l'écran a ses libellés).
# `noms` : ce qu'une feuille ou un fichier de cette nature a de chances de
# s'appeler — de quoi départager deux natures aux colonnes semblables.
TYPES = {
    "roles": {
        "noms": ["role", "roles", "poste", "postes", "fonction", "fonctions", "position",
                 "positions", "job", "jobs", "metier", "metiers"],
        "champs": [
            {"cle": "nom", "requis": True, "aide": "intitulé du rôle ou du poste",
             "syn": ["nom", "name", "role", "rôle", "intitule", "intitulé", "poste", "fonction",
                     "title", "job", "job title", "position", "metier", "métier"]},
            {"cle": "mission", "requis": False, "aide": "mission ou description du rôle",
             "syn": ["mission", "missions", "mission generale", "description", "descriptif",
                     "objectif", "purpose", "responsabilites", "responsabilités",
                     "responsibilities"]},
        ],
    },
    "taches": {
        "noms": ["tache", "taches", "task", "tasks", "activite", "activites", "activity",
                 "activities", "processus", "process"],
        "champs": [
            {"cle": "activite", "requis": True,
             "aide": "nom de l'activité à laquelle la tâche appartient",
             "syn": ["activite", "activité", "activity", "processus", "process",
                     "semi finish", "semi-finish", "nom de l'activite"]},
            {"cle": "tache", "requis": True, "aide": "intitulé de la tâche",
             "syn": ["tache", "tâche", "task", "etape", "étape", "step", "operation",
                     "opération"]},
            {"cle": "description", "requis": False, "aide": "description ou commentaire de la tâche",
             "syn": ["description", "commentaire", "comment", "commentary", "detail", "détail",
                     "note", "notes"]},
            {"cle": "outils", "requis": False, "aide": "outils utilisés, séparés par des virgules",
             "syn": ["outil", "outils", "tool", "tools", "logiciel", "logiciels", "software",
                     "equipement", "équipement", "equipment"]},
            {"cle": "garant", "requis": False, "aide": "rôle garant de l'activité",
             "syn": ["garant", "guarantor", "responsable", "owner", "accountable"]},
            {"cle": "realisateur", "requis": False, "aide": "rôle qui réalise la tâche",
             "syn": ["realisateur", "réalisateur", "doer", "executant", "exécutant", "executor",
                     "acteur", "responsible", "realise par"]},
            {"cle": "approbateur", "requis": False, "aide": "rôle qui valide la tâche",
             "syn": ["approbateur", "approver", "checker", "valideur", "validateur",
                     "valide par"]},
            {"cle": "competences", "requis": False,
             "aide": "compétences requises, séparées par des virgules",
             "syn": ["competence", "compétence", "competences", "compétences", "skills", "skill",
                     "savoir", "savoirs", "knowledge"]},
        ],
    },
    "outils": {
        "noms": ["outil", "outils", "tool", "tools", "logiciel", "logiciels", "software",
                 "equipement", "equipements", "equipment", "application", "applications"],
        "champs": [
            {"cle": "nom", "requis": True, "aide": "nom de l'outil, du logiciel ou de l'équipement",
             "syn": ["nom", "name", "outil", "tool", "logiciel", "software", "equipement",
                     "équipement", "equipment", "application"]},
            {"cle": "description", "requis": False, "aide": "description ou usage de l'outil",
             "syn": ["description", "usage", "utilisation", "commentaire", "comment", "detail",
                     "détail", "notes"]},
        ],
    },
}

# Les colonnes d'une liste de COLLABORATEURS : repérée pour être écartée,
# jamais importée ici.
COMPTES = {
    "prenom": ["prenom", "prénom", "first name", "firstname", "given name", "forename"],
    "nom": ["nom", "nom de famille", "last name", "lastname", "surname", "family name"],
    "email": ["email", "e-mail", "mail", "courriel", "adresse mail", "adresse e-mail",
              "email address", "e mail", "adresse electronique"],
}

# Une colonne qui porte le nom ENTIER d'une personne.
NOM_COMPLET = {"nom complet", "full name", "fullname", "nom et prenom", "nom prenom",
               "prenom nom", "prenom et nom", "collaborateur", "employee", "employe",
               "salarie", "personne", "person", "identite", "name surname", "complete name"}

EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _lang():
    return session.get("lang", "fr")


def _norm(v):
    """Minuscules, sans accents, tout ce qui n'est ni lettre ni chiffre → espace."""
    s = unicodedata.normalize("NFKD", str(v or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s).split())


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ══════════════════════════════════════════════════════════════════════
#  Lire un fichier
# ══════════════════════════════════════════════════════════════════════
class FichierIllisible(Exception):
    pass


def _lire_feuilles(fichier):
    """[(nom de feuille, lignes)] — chaque ligne est une liste de textes.

    xlsx / xlsm par openpyxl (valeurs calculées, jamais les formules), csv avec
    son séparateur deviné ; la « feuille » d'un csv porte le nom du fichier,
    qui dit souvent ce qu'il contient. ⚠️ `.xls` (l'ancien format binaire)
    n'est pas lu : openpyxl ne le sait pas, et l'annoncer comme accepté ferait
    échouer plus loin avec un message obscur.
    """
    nom = fichier.filename or ""
    if not nom.lower().endswith(EXTENSIONS):
        raise FichierIllisible(t("imph.err_format"))
    donnees = fichier.read(MAX_OCTETS + 1)
    if len(donnees) > MAX_OCTETS:
        raise FichierIllisible(t("imph.err_trop_gros"))
    if not donnees:
        raise FichierIllisible(t("imph.err_vide"))

    if nom.lower().endswith(".csv"):
        texte = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texte = donnees.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            sep = csv.Sniffer().sniff(texte[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ";" if texte.count(";") > texte.count(",") else ","
        lignes = [[_txt(c) for c in l] for l in csv.reader(io.StringIO(texte), delimiter=sep)]
        return [(os.path.splitext(os.path.basename(nom))[0], _elaguer(lignes))]

    try:
        wb = openpyxl.load_workbook(io.BytesIO(donnees), data_only=True, read_only=True)
    except Exception:
        raise FichierIllisible(t("imph.err_illisible"))
    feuilles = []
    for ws in wb.worksheets:
        lignes = []
        for row in ws.iter_rows(values_only=True):
            lignes.append([_txt(c) for c in row])
            if len(lignes) > MAX_LIGNES + 50:
                break
        feuilles.append((ws.title, _elaguer(lignes)))
    wb.close()
    return feuilles


def _elaguer(lignes):
    """Retire les cellules vides en fin de ligne et les lignes vides en fin de feuille."""
    out = []
    for l in lignes:
        while l and not l[-1]:
            l = l[:-1]
        out.append(l)
    while out and not out[-1]:
        out.pop()
    return out


# ══════════════════════════════════════════════════════════════════════
#  Reconnaître les colonnes
# ══════════════════════════════════════════════════════════════════════
def _synonymes(type_, champ):
    """Les synonymes du champ, plus ses deux libellés d'écran (FR et EN) : un
    modèle téléchargé puis rempli doit toujours être reconnu."""
    syn = {_norm(s) for s in champ["syn"]}
    for lang in ("fr", "en"):
        syn.add(_norm(t("imph.f_%s_%s" % (type_, champ["cle"]), lang)))
    return {s for s in syn if s}


def _score(entete, synonymes):
    """3 = identique ; 2 = un synonyme apparaît comme MOTS entiers de l'en-tête.
    Au mot près : « prenom » ne contient pas le mot « nom »."""
    e = _norm(entete)
    if not e:
        return 0
    if e in synonymes:
        return 3
    mots = f" {e} "
    return 2 if any(f" {s} " in mots for s in synonymes) else 0


def _correspondance(type_, entetes):
    """{champ: index de colonne} pour une ligne d'en-têtes. Une colonne ne sert
    qu'à un champ ; les correspondances exactes passent avant les partielles."""
    candidats = []
    for c in TYPES[type_]["champs"]:
        syn = _synonymes(type_, c)
        for i, e in enumerate(entetes):
            sc = _score(e, syn)
            if sc:
                candidats.append((-sc, i, c["cle"]))
    candidats.sort()
    prises, cols = {}, set()
    for _, i, cle in candidats:
        if cle not in prises and i not in cols:
            prises[cle] = i
            cols.add(i)
    return prises


def _colonne_emails(donnees):
    """Une colonne dont les valeurs sont des adresses e-mail, quel que soit son titre."""
    largeur = max((len(l) for l in donnees), default=0)
    for j in range(largeur):
        vals = [l[j] for l in donnees if j < len(l) and l[j]]
        if len(vals) >= 2 and sum(1 for v in vals if EMAIL.match(v)) >= 0.6 * len(vals):
            return True
    return False


def _personnes(lignes):
    """Cette feuille est-elle une liste de COLLABORATEURS ?

    Un nom de famille à côté d'un prénom, ou d'une colonne d'e-mails. ⚠️ Le nom
    doit être un titre EXACT (« Nom », « Last name ») : « Nom du rôle » ne
    désigne pas une personne, et un rôle au nom composé (« Chef de projet »)
    n'est pas un nom complet — c'est le titre de colonne qui tranche.
    """
    syn = {k: {_norm(s) for s in v} for k, v in COMPTES.items()}
    for idx, ligne in enumerate(lignes[:10]):
        nom = any(_score(e, syn["nom"]) == 3 or _norm(e) in NOM_COMPLET for e in ligne)
        if not nom:
            continue
        prenom = any(_score(e, syn["prenom"]) for e in ligne)
        email = any(_score(e, syn["email"]) for e in ligne) or _colonne_emails(
            lignes[idx + 1: idx + 40])
        if prenom or email:
            return True
    return False


def _nom_correspond(type_, nom):
    mots = f" {_norm(nom)} "
    return any(f" {s} " in mots for s in TYPES[type_]["noms"])


def _evaluer(type_, nom_feuille, lignes):
    """La meilleure ligne d'en-tête de cette feuille pour cette nature :
    (score, index, correspondance) ou None."""
    requis = [c["cle"] for c in TYPES[type_]["champs"] if c["requis"]]
    nom_ok = _nom_correspond(type_, nom_feuille)
    meilleur = None
    for idx, ligne in enumerate(lignes[:10]):
        cor = _correspondance(type_, ligne)
        couverts = set(cor)
        if not couverts:
            continue
        n_req = sum(1 for r in requis if r in couverts)
        score = (n_req == len(requis), nom_ok, n_req, len(couverts))
        if meilleur is None or score > meilleur[0]:
            meilleur = (score, idx, cor)
    return meilleur


def _nettoyer_liste(valeur):
    """« SAP ; Excel, - » → « SAP, Excel ». Les mentions d'absence (« No special
    skills required », « - », « n/a ») disparaissent DÈS la lecture : l'aperçu
    montre exactement ce qui sera importé."""
    from Code.routes.import_full import est_mention_vide
    return ", ".join(x.strip() for x in re.split(r"[,;\n]", valeur or "")
                     if not est_mention_vide(x))


def _propager_taches(lignes):
    """Cellules fusionnées et lignes de section : une activité (et son garant)
    vaut pour les lignes qui suivent jusqu'à la suivante. Une ligne sans tâche
    n'est pas importée ; si elle ne porte que des outils, ils rejoignent la
    tâche du dessus (tableaux où les outils d'une tâche s'étalent).

    ⚠️ Le garant ne passe PAS d'une activité à l'autre : une activité sans
    garant écrit héritait de celui du bloc précédent, et l'import aurait lié
    ce rôle, à tort, à une activité qu'il ne tient pas. Mieux vaut un garant
    manquant — visible, à compléter — qu'un garant faux."""
    out, dernier = [], {"activite": "", "garant": ""}
    for d in lignes:
        d["outils"] = _nettoyer_liste(d.get("outils"))
        d["competences"] = _nettoyer_liste(d.get("competences"))
        if d.get("activite") and _norm(d["activite"]) != _norm(dernier["activite"]):
            dernier = {"activite": d["activite"], "garant": ""}
        for k in ("activite", "garant"):
            if d.get(k):
                dernier[k] = d[k]
            else:
                d[k] = dernier[k]
        if not d.get("tache"):
            if d.get("outils") and out and out[-1]["activite"] == d["activite"]:
                out[-1]["outils"] = ", ".join(x for x in (out[-1]["outils"], d["outils"]) if x)
            continue
        out.append(d)
    return out


def _extraire(type_, lignes, idx, cor):
    """Les lignes de données en dictionnaires {champ: texte}, un index `_i`
    pour que l'écran s'y retrouve."""
    champs = [c["cle"] for c in TYPES[type_]["champs"]]
    out = []
    for ligne in lignes[idx + 1:]:
        d = {c: "" for c in champs}
        d.update({cle: (ligne[i] if i < len(ligne) else "") for cle, i in cor.items()})
        if any(d.values()):
            out.append(d)
    if type_ == "taches":
        out = _propager_taches(out)
    for i, d in enumerate(out):
        d["_i"] = i
    return out


def _part(pid, fichier, feuille, type_, lignes, idx=None, cor=None, source="FICHIER", **extra):
    """Ce que l'écran reçoit pour une feuille : reconnue ou non, et pourquoi."""
    cor = cor or {}
    rep = {
        "id": pid, "fichier": fichier, "feuille": feuille, "type": type_,
        "source": source, "ligne_entete": idx,
        "entetes": lignes[idx] if idx is not None and 0 <= idx < len(lignes) else [],
        "correspondance": cor,
        "echantillon": [l[:12] for l in lignes[:7]],
        "reconnu": False, "manquants": [], "lignes": [],
    }
    rep.update(extra)
    if type_ is None or type_ == "comptes":
        return rep
    rep["manquants"] = [c["cle"] for c in TYPES[type_]["champs"]
                        if c["requis"] and c["cle"] not in cor]
    if rep["manquants"] or idx is None:
        return rep
    donnees = _extraire(type_, lignes, idx, cor)
    if len(donnees) > MAX_LIGNES:
        rep["erreur"] = t("imph.err_trop_de_lignes").replace("{n}", str(MAX_LIGNES))
        return rep
    rep.update(reconnu=True, lignes=donnees)
    return rep


def _analyser_feuille(pid, fichier, nom, lignes, types, force=False):
    """La nature d'une feuille parmi `types`, sa ligne d'en-tête et ses colonnes.

    Deux natures à égalité (« Nom | Description » : un rôle ou un outil ?) :
    on garde la première dans `ORDRE`, et l'écran demande de trancher.
    ⚠️ Une liste de COLLABORATEURS est écartée (`comptes`) au lieu d'être lue
    comme la nature demandée — les noms de famille deviendraient des rôles.
    Seul un tableau de tâches COMPLET (activité + tâche) passe outre, et
    l'utilisateur, s'il dit lui-même ce que contient la feuille (`force`).
    """
    notes = []
    for ty in types:
        r = _evaluer(ty, nom, lignes)
        if r:
            notes.append((r[0], ty, r))
    meilleur = max((sc for sc, _, _ in notes), default=None)
    ex_aequo = sorted((ty for sc, ty, _ in notes if sc == meilleur), key=ORDRE.index)
    taches_completes = bool(meilleur and meilleur[0] and "taches" in ex_aequo)
    if not force and not taches_completes and _personnes(lignes):
        return _part(pid, fichier, nom, "comptes", lignes)
    if not notes:
        return _part(pid, fichier, nom, types[0] if len(types) == 1 else None, lignes)
    _, idx, cor = next(r for sc, ty, r in notes if ty == ex_aequo[0])
    return _part(pid, fichier, nom, ex_aequo[0], lignes, idx, cor,
                 ambigu=ex_aequo if len(ex_aequo) > 1 else None)


# ══════════════════════════════════════════════════════════════════════
#  L'IA désigne les colonnes, elle n'écrit rien
# ══════════════════════════════════════════════════════════════════════
def _appeler_ia(systeme, contenu):
    """Un appel JSON au fournisseur configuré ; None s'il n'y en a pas ou s'il
    ne répond pas. Isolé : les tests le remplacent."""
    from Code.routes.propose_common import ai_model, openai_client_or_none
    client, _err = openai_client_or_none()
    if client is None:
        return None
    try:
        resp = client.chat.completions.create(
            model=ai_model(),
            messages=[{"role": "system", "content": systeme},
                      {"role": "user", "content": json.dumps(contenu, ensure_ascii=False)}],
            temperature=0, response_format={"type": "json_object"})
        return json.loads(resp.choices[0].message.content or "{}")
    except Exception:
        current_app.logger.exception("import : appel IA")
        return None


def _ia_disponible():
    from Code.prompts import get_prompt
    from Code.routes.propose_common import openai_client_or_none
    client, _err = openai_client_or_none()
    return client is not None and get_prompt("import.correspondance") is not None


def _organiser_ia(pid, fichier, nom, lignes, types):
    """La feuille relue avec les colonnes que l'IA désigne — ou None."""
    from Code.prompts import get_prompt
    systeme = get_prompt("import.correspondance")
    if systeme is None:
        return None
    echantillon = lignes[:20]
    contenu = {
        "langue_des_remarques": "français" if _lang() == "fr" else "English",
        "nom_de_la_feuille": nom,
        "schemas": {ty: [{"champ": c["cle"], "contenu": c["aide"], "obligatoire": c["requis"]}
                         for c in TYPES[ty]["champs"]] for ty in types},
        "lignes": [{"ligne": i, "cellules": {str(j): v[:80] for j, v in enumerate(l[:30]) if v}}
                   for i, l in enumerate(echantillon)],
    }
    brut = _appeler_ia(systeme, contenu)
    if not isinstance(brut, dict):
        return None

    remarque = str(brut.get("remarque") or "")[:300]
    ty = brut.get("type")
    if len(types) == 1 and ty is None and brut.get("colonnes"):
        ty = types[0]
    if ty not in types:
        return _part(pid, fichier, nom, types[0] if len(types) == 1 else None, lignes,
                     source="IA", remarque=remarque, confiance="low")

    largeur = max((len(l) for l in lignes[:60]), default=0)
    try:
        idx = int(brut.get("ligne_entete", 0))
    except (TypeError, ValueError):
        idx = 0
    idx = max(-1, min(idx, len(echantillon) - 1))
    cles = {c["cle"] for c in TYPES[ty]["champs"]}
    cor, prises = {}, set()
    for cle, col in (brut.get("colonnes") or {}).items():
        if isinstance(col, str) and col.isdigit():
            col = int(col)
        if cle in cles and isinstance(col, int) and 0 <= col < largeur and col not in prises:
            cor[cle] = col
            prises.add(col)
    conf = brut.get("confiance") if brut.get("confiance") in ("high", "medium", "low") else "medium"
    return _part(pid, fichier, nom, ty, lignes, idx, cor, source="IA",
                 remarque=remarque, confiance=conf)


# ══════════════════════════════════════════════════════════════════════
#  Qui importe quoi, et où
# ══════════════════════════════════════════════════════════════════════
def _moi():
    from Code.permissions import current_user
    return current_user()


def _cibles_possibles(moi):
    """Les cartos où ce compte écrit : c'est là, et seulement là, qu'il importe."""
    from Code.carto_access import can_edit
    return [e for e in Entity.accessible(moi.id) if can_edit(e, moi)] if moi else []


def _cibles_demandees(moi, ids):
    """Les cartos demandées, restreintes à celles où ce compte écrit — un id
    venu du navigateur ne suffit jamais à écrire dans une carto."""
    permises = {e.id: e for e in _cibles_possibles(moi)}
    out = []
    for i in ids or []:
        try:
            e = permises.get(int(i))
        except (TypeError, ValueError):
            continue
        if e is not None and e not in out:
            out.append(e)
    return out


def _peut_importer(moi):
    return bool(moi) and bool(_cibles_possibles(moi))


def _lignes(type_, brutes):
    """Les lignes renvoyées par le navigateur, ramenées aux champs connus."""
    cles = [c["cle"] for c in TYPES[type_]["champs"]]
    out = []
    for b in (brutes or [])[:MAX_LIGNES]:
        if not isinstance(b, dict):
            continue
        d = {c: str(b.get(c) or "").strip()[:4000 if c in ("description", "mission") else 500]
             for c in cles}
        try:
            d["_i"] = int(b.get("_i", len(out)))
        except (TypeError, ValueError):
            d["_i"] = len(out)
        out.append(d)
    return out


# ══════════════════════════════════════════════════════════════════════
#  Vérifier : le statut de chaque ligne selon la portée
# ══════════════════════════════════════════════════════════════════════
def _statut(n_cibles, n_deja):
    if n_deja == 0:
        return "nouveau"
    return "present" if n_deja >= n_cibles else "partiel"


def _noms_existants(modele, entity_id):
    """Les noms déjà posés dans une carto, sous leur forme normalisée — et pour
    un rôle, ses traductions : « Quality » EST le rôle « Qualité »."""
    noms = set()
    for x in modele.query.filter_by(entity_id=entity_id).all():
        noms.add(_norm(x.name))
        for attr in ("name_fr", "name_en"):
            if getattr(x, attr, None):
                noms.add(_norm(getattr(x, attr)))
    return noms


def _verifier_simple(type_, lignes, cibles):
    """Rôles et outils : un nom par ligne, créé dans chaque carto où il manque."""
    modele = Role if type_ == "roles" else Tool
    existants = {e.id: _noms_existants(modele, e.id) for e in cibles}
    vus, out = set(), []
    for l in lignes:
        nom = l.get("nom", "")
        r = dict(l)
        if not nom:
            r.update(statut="invalide", raison=t("imph.st_nom_vide"))
        elif _norm(nom) in vus:
            r.update(statut="invalide", raison=t("imph.st_doublon"))
        elif len(nom) > (100 if type_ == "roles" else 255):
            r.update(statut="invalide", raison=t("imph.st_trop_long"))
        else:
            deja = sum(1 for e in cibles if _norm(nom) in existants[e.id])
            r.update(statut=_statut(len(cibles), deja), n_deja=deja,
                     n_nouveau=len(cibles) - deja)
        vus.add(_norm(nom))
        out.append(r)
    return {"lignes": out}


def _activites(cibles):
    """{nom normalisé: {"nom": …, "cartos": {entity_id: activité}}} — l'union
    des activités des cartos visées. Une tâche va dans chaque carto où son
    activité existe sous ce nom."""
    acts = {}
    for e in cibles:
        for a in Activities.query.filter_by(entity_id=e.id).order_by(Activities.name).all():
            if not a.name:
                continue
            acts.setdefault(_norm(a.name), {"nom": a.name, "cartos": {}})["cartos"][e.id] = a
    return acts


def _proches(nom, acts, n=3):
    """Les activités les plus proches d'un nom du fichier : [(score, clé)]."""
    cle = _norm(nom)
    notes = []
    for k in acts:
        s = SequenceMatcher(None, cle, k).ratio()
        if cle and (cle in k or k in cle):
            s = max(s, 0.88)
        notes.append((round(s, 2), k))
    notes.sort(key=lambda x: (-x[0], x[1]))
    return notes[:n]


def _groupes(lignes):
    """Les lignes regroupées par activité, dans l'ordre du fichier."""
    groupes, par_cle = [], {}
    for l in lignes:
        cle = _norm(l.get("activite"))
        g = par_cle.get(cle)
        if g is None:
            g = {"cle": cle, "activite_fichier": l.get("activite", ""), "garant": "",
                 "lignes": []}
            par_cle[cle] = g
            groupes.append(g)
        if not g["garant"] and l.get("garant"):
            g["garant"] = l["garant"]
        g["lignes"].append(l)
    return groupes


def _verifier_taches(lignes, cibles, choix):
    """Les tâches regroupées par activité du fichier, chaque groupe rattaché à
    une activité des cartos visées.

    Rattachement automatique seulement au nom près, ou à 90 % de ressemblance ;
    en dessous l'écran propose et l'utilisateur choisit. ⚠️ `choix` l'emporte
    toujours : une chaîne vide veut dire « ignorer ce groupe ».
    """
    acts = _activites(cibles)
    groupes = _groupes(lignes)
    choix = {_norm(k): v for k, v in (choix or {}).items()}
    rattachees = {}
    for g in groupes:
        proches = _proches(g["activite_fichier"], acts)
        g["possibles"] = [{"nom": acts[k]["nom"], "score": s} for s, k in proches]
        if not g["cle"]:
            g["mode"], g["choix"] = "sans_activite", None
        elif g["cle"] in choix:
            cle = _norm(choix[g["cle"]])
            g["mode"], g["choix"] = ("manuel", acts[cle]["nom"]) if cle in acts else ("ignore", None)
        elif g["cle"] in acts:
            g["mode"], g["choix"] = "exact", acts[g["cle"]]["nom"]
        elif proches and proches[0][0] >= 0.9:
            g["mode"], g["choix"] = "proche", acts[proches[0][1]]["nom"]
        else:
            g["mode"], g["choix"] = "a_rattacher", None
        if g["choix"]:
            rattachees[_norm(g["choix"])] = acts[_norm(g["choix"])]["cartos"]

    ids = [a.id for cartos in rattachees.values() for a in cartos.values()]
    deja = {}
    for tache in (Task.query.filter(Task.activity_id.in_(ids)).all() if ids else []):
        deja.setdefault(tache.activity_id, set()).add(_norm(tache.name))

    for g in groupes:
        cartos = rattachees.get(_norm(g["choix"])) if g["choix"] else None
        g["n_cartos"] = len(cartos or {})
        g["n_cibles"] = len(cibles)
        vues = set()
        for l in g["lignes"]:
            nom = _norm(l.get("tache"))
            if g["mode"] == "sans_activite":
                l.update(statut="invalide", raison=t("imph.st_sans_activite"))
            elif nom in vues:
                l.update(statut="invalide", raison=t("imph.st_doublon"))
            elif len(l.get("tache", "")) > 255:
                l.update(statut="invalide", raison=t("imph.st_trop_long"))
            elif g["mode"] == "ignore":
                l.update(statut="ignore")
            elif not cartos:
                l.update(statut="a_rattacher")
            else:
                n = sum(1 for a in cartos.values() if nom in deja.get(a.id, ()))
                l.update(statut=_statut(len(cartos), n), n_deja=n, n_nouveau=len(cartos) - n)
            vues.add(nom)
    return {"groupes": groupes, "lignes": [l for g in groupes for l in g["lignes"]],
            "activites": sorted((v["nom"] for v in acts.values()), key=lambda s: s.lower())}


def _totaux(lignes):
    tot = {"nouveau": 0, "partiel": 0, "present": 0, "invalide": 0, "a_rattacher": 0,
           "ignore": 0}
    for l in lignes:
        s = l.get("statut", "invalide")
        tot[s] = tot.get(s, 0) + 1
    return tot


# ══════════════════════════════════════════════════════════════════════
#  Importer
# ══════════════════════════════════════════════════════════════════════
def _importer_simple(type_, lignes, cibles):
    """Rôles et outils : créés dans chaque carto visée où ils manquent. Un rôle
    importé n'est pas une bande : `hors_carte`, sinon le prochain
    enregistrement de la carte l'effacerait."""
    from Code.role_i18n import on_role_name_saved
    verifiees = _verifier_simple(type_, lignes, cibles)["lignes"]
    modele = Role if type_ == "roles" else Tool
    existants = {e.id: _noms_existants(modele, e.id) for e in cibles}
    crees, par_carto = 0, {}
    for l in verifiees:
        if l["statut"] not in ("nouveau", "partiel"):
            continue
        for e in cibles:
            if _norm(l["nom"]) in existants[e.id]:
                continue
            if type_ == "roles":
                obj = Role(name=l["nom"], entity_id=e.id, hors_carte=True,
                           mission_generale=l.get("mission") or None)
                on_role_name_saved(obj, l["nom"])
            else:
                obj = Tool(name=l["nom"], entity_id=e.id, description=l.get("description") or None)
            db.session.add(obj)
            existants[e.id].add(_norm(l["nom"]))
            crees += 1
            par_carto[e.name] = par_carto.get(e.name, 0) + 1
    db.session.flush()
    return {"crees": crees, "par_carto": par_carto}


def _liste(v):
    return [x.strip() for x in re.split(r"[,;\n]", v or "") if x.strip()]


def _importer_taches(lignes, cibles, choix):
    """Chaque groupe rattaché passe par le pipeline d'import de l'application
    (`injecter_groupes`), carto par carto : mêmes get-or-create d'outils et de
    rôles, mêmes liens garant / réalisateur / approbateur, même déduplication."""
    from Code.routes.import_full import injecter_groupes
    verif = _verifier_taches(lignes, cibles, choix)
    acts = _activites(cibles)
    stats = {"tasks_created": 0, "tools_created": 0, "roles_created": 0,
             "competencies_created": 0, "activities_updated": 0}
    for g in verif["groupes"]:
        if not g["choix"]:
            continue
        taches = [{"name": l["tache"], "tools": _liste(l.get("outils")),
                   "doer": l.get("realisateur", ""), "approver": l.get("approbateur", ""),
                   "skills": _liste(l.get("competences")), "commentary": l.get("description", "")}
                  for l in g["lignes"] if l.get("statut") in ("nouveau", "partiel")]
        if not taches:
            continue
        for eid, a in acts[_norm(g["choix"])]["cartos"].items():
            s = injecter_groupes([{"activity_id": a.id, "guarantor": g["garant"],
                                   "tasks": taches}], eid)
            for k in stats:
                stats[k] += s.get(k, 0)
    db.session.flush()
    return stats


# ══════════════════════════════════════════════════════════════════════
#  Routes
# ══════════════════════════════════════════════════════════════════════
@import_hub_bp.route("/contexte", methods=["GET"])
def contexte():
    """Ce que la fenêtre doit savoir avant tout : les cartos où écrire, et si
    l'IA est là. Et où s'importent les comptes — pas ici."""
    from Code.permissions import can_create_accounts
    moi = _moi()
    if moi is None:
        return jsonify({"error": "unauthorized"}), 401
    active = Entity.get_active_id()
    cibles = _cibles_possibles(moi)
    natures = {k: {"nom": t("imph.t_%s" % k), "desc": t("imph.d_%s" % k),
                   "portee": t("imph.p_%s" % k)} for k in list(TYPES) + ["multiple"]}
    for k, v in TYPES.items():
        natures[k]["champs"] = [{"cle": c["cle"], "requis": c["requis"],
                                 "label": t("imph.f_%s_%s" % (k, c["cle"]))} for c in v["champs"]]
    return jsonify({
        "natures": natures,
        "ordre": list(ORDRE),
        "cibles": [{"id": e.id, "name": e.name, "active": e.id == active,
                    "shared": bool(getattr(e, "is_shared", False))} for e in cibles],
        "active": next(({"id": e.id, "name": e.name} for e in cibles if e.id == active), None),
        "peut": bool(cibles),
        # Les comptes s'importent depuis la page Comptes, et seulement par qui
        # peut en créer : l'écran le dit, avec le lien quand il mène quelque part.
        "comptes": {"url": url_for("gestion_compte.list_users", tab="import-tab"),
                    "peut": bool(can_create_accounts(moi))},
        "ia": _ia_disponible(),
        "max_mo": MAX_OCTETS // (1024 * 1024),
        "max_lignes": MAX_LIGNES,
    }), 200


@import_hub_bp.route("/modele/<type_>", methods=["GET"])
def modele(type_):
    """Un classeur prêt à remplir, avec les en-têtes que la lecture reconnaît.
    Pour l'import multiple : une feuille par nature."""
    if type_ not in TYPES and type_ != "multiple":
        return jsonify({"error": t("imph.err_type")}), 404
    lang = _lang()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    couleurs = {"roles": "059669", "taches": "7C3AED", "outils": "EA580C"}
    for ty in (ORDRE if type_ == "multiple" else (type_,)):
        ws = wb.create_sheet(t("imph.t_%s" % ty, lang)[:31])
        champs = TYPES[ty]["champs"]
        ws.append([t("imph.f_%s_%s" % (ty, c["cle"]), lang) for c in champs])
        for cell, c in zip(ws[1], champs):
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF", underline=(
                "single" if c["requis"] else None))
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor=couleurs[ty])
        for i in range(1, len(champs) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 30
        ws.freeze_panes = "A2"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(tampon, as_attachment=True,
                     download_name="%s-%s.xlsx" % (
                         t("imph.modele_fichier", lang),
                         _norm(t("imph.t_%s" % type_, lang)).replace(" ", "-")),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _types_voulus(type_, comme):
    if comme in TYPES:
        return [comme]
    return list(ORDRE) if type_ == "multiple" else [type_]


def _type_valide(type_):
    return type_ in TYPES or type_ == "multiple"


@import_hub_bp.route("/lire", methods=["POST"])
def lire():
    """Lit le ou les fichiers déposés : une « part » par feuille à importer."""
    moi = _moi()
    type_ = request.form.get("type", "")
    if not _type_valide(type_):
        return jsonify({"error": t("imph.err_type")}), 404
    if not _peut_importer(moi):
        return jsonify({"error": t("imph.err_droits")}), 403
    fichiers = request.files.getlist("fichiers")[:MAX_FICHIERS if type_ == "multiple" else 1]
    if not fichiers:
        return jsonify({"error": t("imph.err_vide")}), 400
    feuille_voulue = request.form.get("feuille")
    comme = request.form.get("comme")
    types = _types_voulus(type_, comme)
    # Quand l'utilisateur dit lui-même ce que contient la feuille, on le croit :
    # c'est la seule façon de passer outre un repérage de collaborateurs erroné.
    force = comme in TYPES

    parts, erreurs = [], []
    for i, f in enumerate(fichiers):
        try:
            feuilles = _lire_feuilles(f)
        except FichierIllisible as e:
            erreurs.append({"fichier": f.filename, "erreur": str(e)})
            continue
        feuilles = [(n, l) for n, l in feuilles if l and (not feuille_voulue or n == feuille_voulue)]
        if not feuilles:
            erreurs.append({"fichier": f.filename, "erreur": t("imph.err_vide")})
            continue
        lues = [_analyser_feuille("%d:%s" % (i, n), f.filename, n, l, types, force)
                for n, l in feuilles]
        for p in lues:
            p["n_feuilles"] = len(feuilles)
        if type_ == "multiple":
            parts.extend(lues)
        else:
            # Une seule nature attendue : la feuille qui la porte le mieux.
            parts.append(max(lues, key=lambda p: (p["reconnu"], len(p["lignes"]),
                                                  len(p["correspondance"]))))
    return jsonify({"parts": parts, "erreurs": erreurs}), 200


@import_hub_bp.route("/organiser", methods=["POST"])
def organiser():
    """Une feuille que la lecture n'a pas reconnue, relue avec les colonnes que
    l'IA désigne. Le fichier est renvoyé : rien n'est gardé côté serveur.
    Rien n'est appliqué non plus : l'écran montre cette lecture et attend
    qu'on l'accepte."""
    moi = _moi()
    type_ = request.form.get("type", "")
    if not _type_valide(type_):
        return jsonify({"error": t("imph.err_type")}), 404
    if not _peut_importer(moi):
        return jsonify({"error": t("imph.err_droits")}), 403
    f = request.files.get("fichier")
    if f is None:
        return jsonify({"error": t("imph.err_vide")}), 400
    try:
        feuilles = [(n, l) for n, l in _lire_feuilles(f) if l]
    except FichierIllisible as e:
        return jsonify({"error": str(e)}), 400
    voulue = request.form.get("feuille")
    feuille = next(((n, l) for n, l in feuilles if n == voulue), None) or (
        max(feuilles, key=lambda x: len(x[1])) if feuilles else None)
    if feuille is None:
        return jsonify({"error": t("imph.err_vide")}), 400
    nom, lignes = feuille
    part = _organiser_ia(request.form.get("id") or "0:" + nom, f.filename, nom, lignes,
                         _types_voulus(type_, request.form.get("comme")))
    if part is None:
        return jsonify({"error": t("imph.err_ia")}), 503
    part["n_feuilles"] = len(feuilles)
    return jsonify({"part": part}), 200


@import_hub_bp.route("/verifier", methods=["POST"])
def verifier():
    """Le statut de chaque ligne contre la base, pour la portée choisie. N'écrit rien."""
    moi = _moi()
    p = request.get_json(silent=True) or {}
    type_ = p.get("type")
    if type_ not in TYPES:
        return jsonify({"error": t("imph.err_type")}), 404
    if not _peut_importer(moi):
        return jsonify({"error": t("imph.err_droits")}), 403
    cibles = _cibles_demandees(moi, p.get("cibles"))
    if not cibles:
        return jsonify({"error": t("imph.err_aucune_cible")}), 400
    lignes = _lignes(type_, p.get("lignes"))
    res = (_verifier_taches(lignes, cibles, p.get("choix") or {}) if type_ == "taches"
           else _verifier_simple(type_, lignes, cibles))
    res["totaux"] = _totaux(res["lignes"])
    res["cibles"] = [e.id for e in cibles]
    return jsonify(res), 200


@import_hub_bp.route("/rapprocher", methods=["POST"])
def rapprocher():
    """L'IA propose, pour les activités du fichier restées sans correspondance,
    l'activité des cartos visées qui leur ressemble par le SENS (« Identify
    Part » ↔ « Développer la solution technique »). Rien n'est appliqué :
    l'écran en fait un compte rendu, l'utilisateur garde ce qu'il veut."""
    from Code.prompts import get_prompt
    moi = _moi()
    if not _peut_importer(moi):
        return jsonify({"error": t("imph.err_droits")}), 403
    p = request.get_json(silent=True) or {}
    cibles = _cibles_demandees(moi, p.get("cibles"))
    noms = [str(n)[:300] for n in (p.get("noms") or [])[:80] if str(n).strip()]
    if not cibles or not noms:
        return jsonify({"propositions": {}}), 200
    acts = sorted({v["nom"] for v in _activites(cibles).values()}, key=str.lower)
    systeme = get_prompt("import.enrich")
    # La justification s'affiche telle quelle dans le compte rendu : elle doit
    # être écrite dans la langue de celui qui le lit.
    brut = _appeler_ia(systeme, {
        "unmatched": [{"activity_name_excel": n} for n in noms],
        "db_activities": [{"id": i, "name": n} for i, n in enumerate(acts)],
        "langue_des_remarques": "français" if _lang() == "fr" else "English",
    }) if systeme else None
    if brut is None:
        return jsonify({"error": t("imph.err_ia")}), 503
    propositions = {}
    for r in brut.get("resolved") or []:
        try:
            nom_act = acts[int(r.get("activity_id"))]
        except (TypeError, ValueError, IndexError):
            continue
        if r.get("activity_name_excel") in noms:
            conf = r.get("confidence")
            propositions[r["activity_name_excel"]] = {
                "activite": nom_act,
                "confiance": conf if conf in ("high", "medium", "low") else "medium",
                "raison": str(r.get("match_reason") or "")[:200],
            }
    return jsonify({"propositions": propositions}), 200


@import_hub_bp.route("/importer", methods=["POST"])
def importer():
    """Écrit ce que l'utilisateur a gardé, en UNE transaction : un import
    multiple réussit entièrement ou ne laisse rien derrière lui.

    Chaque part est revérifiée ici — le statut venu du navigateur n'est jamais
    cru — et les natures passent dans l'ordre `ORDRE`.
    """
    moi = _moi()
    p = request.get_json(silent=True) or {}
    parts = p.get("parts") if isinstance(p.get("parts"), list) else [p]
    parts = [x for x in parts if isinstance(x, dict)]
    if not parts or any(x.get("type") not in TYPES for x in parts):
        return jsonify({"error": t("imph.err_type")}), 400
    if not _peut_importer(moi):
        return jsonify({"error": t("imph.err_droits")}), 403
    cibles = _cibles_demandees(moi, p.get("cibles"))
    if not cibles:
        return jsonify({"error": t("imph.err_aucune_cible")}), 400

    resultats = []
    try:
        for x in sorted(parts, key=lambda x: ORDRE.index(x["type"])):
            lignes = _lignes(x["type"], x.get("lignes"))
            if x["type"] == "taches":
                res = _importer_taches(lignes, cibles, x.get("choix") or {})
            else:
                res = _importer_simple(x["type"], lignes, cibles)
            resultats.append(dict(type=x["type"], id=x.get("id"), **res))
        db.session.commit()
    except Exception:
        db.session.rollback()
        current_app.logger.exception("import")
        return jsonify({"error": t("imph.err_import")}), 500
    return jsonify({"ok": True, "resultats": resultats,
                    "cibles": [{"id": e.id, "name": e.name} for e in cibles]}), 200
